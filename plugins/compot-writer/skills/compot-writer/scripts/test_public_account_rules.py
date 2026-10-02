from __future__ import annotations

import copy
import tempfile
import unittest
import zipfile
from pathlib import Path

from docx import Document
from lxml import etree

from build_package import (
    ASSET_NAMES,
    DEFAULT_TEMPLATE,
    NS,
    QN,
    XML,
    enforce_english_quote_font,
    patch_docx,
    validate_public_account_style,
)
from validate_package import validate_quote_fonts


TNR = "Times New Roman"
FONT_ATTRIBUTES = ("ascii", "hAnsi", "eastAsia", "cs")
THEME_ATTRIBUTES = ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme")


def paragraph_text(paragraph: etree._Element) -> str:
    return "".join(paragraph.xpath(".//w:t/text()", namespaces=NS))


def make_paragraph(text: str, *, mixed_content: bool = False) -> etree._Element:
    paragraph = etree.Element(QN("p"), nsmap={"w": NS["w"]})
    ppr = etree.SubElement(paragraph, QN("pPr"))
    etree.SubElement(ppr, QN("spacing"), {QN("after"): "120"})
    run = etree.SubElement(paragraph, QN("r"))
    rpr = etree.SubElement(run, QN("rPr"))
    fonts = etree.SubElement(rpr, QN("rFonts"))
    for name in FONT_ATTRIBUTES:
        fonts.set(QN(name), "宋体")
    for name in THEME_ATTRIBUTES:
        fonts.set(QN(name), "minorEastAsia")
    fonts.set(QN("hint"), "eastAsia")
    etree.SubElement(rpr, QN("b"))
    etree.SubElement(rpr, QN("color"), {QN("val"): "123456"})
    etree.SubElement(rpr, QN("sz"), {QN("val"): "21"})
    if mixed_content:
        drawing = etree.SubElement(run, QN("drawing"))
        etree.SubElement(drawing, "{urn:test}image", {"relationship": "rId8"})
        etree.SubElement(run, QN("br"))
    node = etree.SubElement(run, QN("t"))
    node.set(f"{{{XML}}}space", "preserve")
    node.text = text
    return paragraph


def quote_runs(paragraph: etree._Element) -> list[etree._Element]:
    return [
        run
        for run in paragraph.iter(QN("r"))
        if '"' in "".join(node.text or "" for node in run.findall(QN("t")))
    ]


class PublicAccountStyleTests(unittest.TestCase):
    def test_publication_year_and_issue_narrative_is_rejected(self) -> None:
        for text in (
            "论文于2023年在线发表，收录于2024年第34卷第1期。",
            "该成果于2024年发表。",
            "该论文发表于2023年。",
            "论文收录于第34卷第1期。",
        ):
            with self.subTest(text=text), self.assertRaises(ValueError):
                validate_public_account_style(text)

    def test_source_figure_and_table_locators_are_rejected(self) -> None:
        for text in (
            "图1 HANet整体结构（原文图5）",
            "表1 TinyPerson结果（原文表I）",
            "对应原论文图 5 的整体结构。",
            "原始论文中的图5",
            "原论文中表III给出了详细结果。",
            "实验见原文表 IV。",
            "整体结构（original Fig. 5）",
            "检测结果（original Table I）",
        ):
            with self.subTest(text=text), self.assertRaises(ValueError):
                validate_public_account_style(text)

    def test_current_article_numbering_and_experimental_years_are_allowed(self) -> None:
        for text in (
            "图1 HANet整体结构",
            "表1 TinyPerson上的检测结果",
            "2023年采集的数据用于测试，2024年的样本用于训练。",
            "原文表明，该结构改善了极小目标的检测。",
            "如图1所示，实验结果见表1。",
            '该成果以"Paper Title"为题，发表在"Journal Name"上。',
            "original figure legend",
            '该成果以"The original figure legend"为题，发表在"Journal Name"上。',
        ):
            with self.subTest(text=text):
                validate_public_account_style(text)

    def test_explicit_bibliographic_override_allows_date_narrative(self) -> None:
        validate_public_account_style(
            "论文于2023年在线发表，收录于2024年第34卷第1期。",
            allow_bibliographic_details=True,
        )

    def test_bibliographic_override_does_not_allow_source_locators(self) -> None:
        for text in ("图1（原文图5）", "表1（原文表I）", "结果见原文表 IV。"):
            with self.subTest(text=text), self.assertRaises(ValueError):
                validate_public_account_style(text, allow_bibliographic_details=True)


class EnglishQuoteFontTests(unittest.TestCase):
    def test_quotes_are_standalone_and_explicitly_times_new_roman(self) -> None:
        paragraph = make_paragraph('该成果以"Paper Title"为题，发表在"Journal Name"上。')
        original_text = paragraph_text(paragraph)
        enforce_english_quote_font(paragraph)
        self.assertEqual(paragraph_text(paragraph), original_text)
        self.assertEqual(len(quote_runs(paragraph)), 4)
        for run in quote_runs(paragraph):
            self.assertEqual("".join(run.xpath("./w:t/text()", namespaces=NS)), '"')
            fonts = run.find("w:rPr/w:rFonts", NS)
            self.assertIsNotNone(fonts)
            for name in FONT_ATTRIBUTES:
                self.assertEqual(fonts.get(QN(name)), TNR)
            self.assertEqual(fonts.get(QN("hint")), "default")
            for name in THEME_ATTRIBUTES:
                self.assertIsNone(fonts.get(QN(name)))
        self.assertIsNone(validate_quote_fonts(paragraph))

    def test_body_styles_paragraph_format_and_images_are_preserved(self) -> None:
        paragraph = make_paragraph(' 中文"Paper Title"结尾。 ', mixed_content=True)
        ppr_before = etree.tostring(paragraph.find("w:pPr", NS))
        image_before = etree.tostring(paragraph.find(".//w:drawing", NS))
        run_before = paragraph.find("w:r", NS)
        rpr_before = copy.deepcopy(run_before.find("w:rPr", NS))
        fonts_before = etree.tostring(rpr_before.find("w:rFonts", NS))
        non_font_styles_before = [
            etree.tostring(child) for child in rpr_before if child.tag != QN("rFonts")
        ]
        original_text = paragraph_text(paragraph)
        enforce_english_quote_font(paragraph)
        self.assertEqual(paragraph_text(paragraph), original_text)
        self.assertEqual(etree.tostring(paragraph.find("w:pPr", NS)), ppr_before)
        self.assertEqual(etree.tostring(paragraph.find(".//w:drawing", NS)), image_before)
        self.assertEqual(len(paragraph.xpath(".//w:drawing", namespaces=NS)), 1)
        self.assertEqual(len(paragraph.xpath(".//w:br", namespaces=NS)), 1)
        for run in paragraph.xpath("./w:r", namespaces=NS):
            text = "".join(run.xpath("./w:t/text()", namespaces=NS))
            rpr = run.find("w:rPr", NS)
            self.assertIsNotNone(rpr)
            non_font_styles_after = [
                etree.tostring(child) for child in rpr if child.tag != QN("rFonts")
            ]
            for original_style in non_font_styles_before:
                self.assertIn(original_style, non_font_styles_after)
            if text != '"':
                self.assertEqual(non_font_styles_after, non_font_styles_before)
            if text and text != '"':
                self.assertEqual(etree.tostring(rpr.find("w:rFonts", NS)), fonts_before)
                for node in run.findall("w:t", NS):
                    if node.text and (node.text.startswith(" ") or node.text.endswith(" ")):
                        self.assertEqual(node.get(f"{{{XML}}}space"), "preserve")

    def test_font_enforcement_is_idempotent(self) -> None:
        paragraph = make_paragraph('以"Paper Title"为题。')
        enforce_english_quote_font(paragraph)
        first_xml = etree.tostring(paragraph)
        enforce_english_quote_font(paragraph)
        self.assertEqual(etree.tostring(paragraph), first_xml)

    def test_paragraph_without_ascii_quotes_is_unchanged(self) -> None:
        paragraph = make_paragraph("原文表明，图1和表1支持这一结论。", mixed_content=True)
        original_xml = etree.tostring(paragraph)
        enforce_english_quote_font(paragraph)
        self.assertEqual(etree.tostring(paragraph), original_xml)
        self.assertIsNone(validate_quote_fonts(paragraph))

    def test_mixed_quote_text_is_rejected_even_with_correct_fonts(self) -> None:
        paragraph = make_paragraph('以"Paper Title"为题。')
        fonts = paragraph.find("w:r/w:rPr/w:rFonts", NS)
        for name in FONT_ATTRIBUTES:
            fonts.set(QN(name), TNR)
        for name in THEME_ATTRIBUTES:
            fonts.attrib.pop(QN(name), None)
        fonts.set(QN("hint"), "default")
        with self.assertRaises(ValueError):
            validate_quote_fonts(paragraph)

    def test_wrong_or_inherited_font_is_rejected(self) -> None:
        for name in FONT_ATTRIBUTES:
            for value in ("Arial", None):
                with self.subTest(attribute=name, value=value):
                    paragraph = make_paragraph('"')
                    enforce_english_quote_font(paragraph)
                    fonts = quote_runs(paragraph)[0].find("w:rPr/w:rFonts", NS)
                    if value is None:
                        fonts.attrib.pop(QN(name))
                    else:
                        fonts.set(QN(name), value)
                    with self.assertRaises(ValueError):
                        validate_quote_fonts(paragraph)

    def test_theme_and_east_asia_hint_overrides_are_rejected(self) -> None:
        for name, value in (
            *((name, "minorEastAsia") for name in THEME_ATTRIBUTES),
            ("hint", "eastAsia"),
        ):
            with self.subTest(attribute=name):
                paragraph = make_paragraph('"')
                enforce_english_quote_font(paragraph)
                quote_runs(paragraph)[0].find("w:rPr/w:rFonts", NS).set(QN(name), value)
                with self.assertRaises(ValueError):
                    validate_quote_fonts(paragraph)


class PackageQuoteIntegrationTests(unittest.TestCase):
    def test_real_package_entry_enforces_quotes_without_changing_template_parts(self) -> None:
        sentence = '该成果以"Paper Title"为题，发表在"Journal Name"上。'
        with tempfile.TemporaryDirectory(prefix="compot-quote-test-") as temporary:
            temporary_path = Path(temporary)
            assets_dir = temporary_path / "assets"
            assets_dir.mkdir()
            with zipfile.ZipFile(DEFAULT_TEMPLATE) as source:
                source_members = {name: source.read(name) for name in source.namelist()}
            for index, name in enumerate(ASSET_NAMES, 1):
                (assets_dir / name).write_bytes(source_members[f"word/media/image{index}.png"])
            avatar = temporary_path / "avatar.png"
            avatar.write_bytes(source_members["word/media/image8.png"])
            output = temporary_path / "quote-test.docx"
            patch_docx(
                DEFAULT_TEMPLATE,
                output,
                {3: sentence},
                assets_dir,
                avatar,
                "测试撰稿人",
            )
            document = Document(output)
            self.assertEqual(document.paragraphs[3].text, sentence)
            self.assertEqual(len(document.paragraphs), 36)
            self.assertEqual(len(document.inline_shapes), 8)
            self.assertEqual(len(document.sections), 1)
            for paragraph in document.paragraphs:
                validate_quote_fonts(paragraph._element)
            self.assertEqual(len(quote_runs(document.paragraphs[3]._element)), 4)
            with zipfile.ZipFile(output) as generated:
                generated_members = {name: generated.read(name) for name in generated.namelist()}
            self.assertEqual(set(generated_members), set(source_members))
            for name, source_bytes in source_members.items():
                if name not in ("word/document.xml", "docProps/core.xml"):
                    self.assertEqual(generated_members[name], source_bytes, name)
            source_root = etree.fromstring(source_members["word/document.xml"])
            generated_root = etree.fromstring(generated_members["word/document.xml"])
            source_section = source_root.find("w:body/w:sectPr", NS)
            generated_section = generated_root.find("w:body/w:sectPr", NS)
            self.assertEqual(etree.tostring(generated_section), etree.tostring(source_section))
            source_paragraphs = source_root.findall("w:body/w:p", NS)
            generated_paragraphs = generated_root.findall("w:body/w:p", NS)
            self.assertEqual(len(source_paragraphs), 36)
            for index, source_paragraph in enumerate(source_paragraphs):
                if index == 33:  # The writer card intentionally gains a page break.
                    continue
                source_ppr = source_paragraph.find("w:pPr", NS)
                generated_ppr = generated_paragraphs[index].find("w:pPr", NS)
                source_xml = etree.tostring(source_ppr) if source_ppr is not None else None
                generated_xml = etree.tostring(generated_ppr) if generated_ppr is not None else None
                self.assertEqual(generated_xml, source_xml, f"paragraph {index} style")


if __name__ == "__main__":
    unittest.main()
