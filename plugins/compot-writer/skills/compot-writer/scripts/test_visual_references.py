from __future__ import annotations

import json
import tempfile
import unittest
from pathlib import Path

from docx import Document

from build_package import extract_visual_references, load_manifest, validate_visual_references
from validate_package import validate_document_visual_references


def make_slots() -> dict[int, str]:
    return {
        3: "微小目标的有效信息有限，作者围绕这一问题设计了检测方法。",
        4: '该成果以"Example Paper"为题，发表在"Example Journal"上。',
        6: "图1 总体结构",
        7: "作者将特征处理分成多个阶段，整体结构如图1所示。",
        8: "不同阶段的功能分别对应目标增强和信息筛选。",
        10: "图2 增强模块",
        11: "该模块通过特征交互增强表示，其结构见图2。",
        13: "图3 筛选模块",
        14: "该模块保留有效信息，具体结构如图3所示。",
        16: "图4 融合模块",
        17: "不同阶段的信息在检测端融合，融合过程见图4。",
        19: "表1 定量比较",
        20: "作者对比了不同方法的检测表现，定量结果见表1。",
        22: "图5 定性比较",
        23: "不同方法呈现出不同的漏检情况，定性对比见图5。",
        24: (
            "综上，作者针对微小目标的信息不足问题，设计了分阶段处理、信息筛选和特征融合方法。"
            "定量实验与定性对比共同说明了这一设计的作用，为进一步分析微小目标的有效表示提供了依据，"
            "其适用范围仍需结合实际实验设置理解。"
        ),
    }


def make_manifest() -> dict[str, object]:
    slots = make_slots()
    return {
        "number": 999,
        "title": "测试用方法",
        "author": "测试署名",
        "source_pdf": "test-paper.pdf",
        "lead": [slots[3], slots[4]],
        "figure1": {"caption": slots[6], "text": [slots[7], slots[8]]},
        "figure2": {"caption": slots[10], "text": slots[11]},
        "figure3": {"caption": slots[13], "text": slots[14]},
        "figure4": {"caption": slots[16], "text": slots[17]},
        "table1": {"caption": slots[19], "text": slots[20]},
        "figure5": {"caption": slots[22], "text": slots[23]},
        "conclusion": slots[24],
        "source_link": "https://example.org/test-paper",
    }


def make_document():
    slots = make_slots()
    document = Document()
    for index in range(36):
        document.add_paragraph(slots.get(index, ""))
    return document


class VisualReferenceTests(unittest.TestCase):
    def test_contextual_references_and_caption_labels_are_accepted(self) -> None:
        self.assertIsNone(validate_visual_references(make_slots()))

    def test_captions_alone_do_not_satisfy_body_references(self) -> None:
        slots = make_slots()
        for index in (7, 8, 11, 14, 17, 20, 23):
            slots[index] = "作者分析了这一模块及其检测表现。"
        with self.assertRaisesRegex(ValueError, "Missing body reference to 图1"):
            validate_visual_references(slots)

    def test_each_figure_and_table_requires_its_own_explanation_reference(self) -> None:
        for index, label in ((7, "图1"), (11, "图2"), (14, "图3"),
                             (17, "图4"), (20, "表1"), (23, "图5")):
            with self.subTest(label=label):
                slots = make_slots()
                slots[index] = "作者分析了模块工作方式与实验表现。"
                # An opening list cannot compensate for an unreferenced explanation.
                slots[3] += "各部分内容分别见图1、2、3、4、5和表1。"
                with self.assertRaisesRegex(ValueError, f"Missing body reference to {label}"):
                    validate_visual_references(slots)

    def test_either_figure_one_paragraph_may_supply_its_reference(self) -> None:
        slots = make_slots()
        slots[7] = "作者将特征处理分成多个阶段。"
        slots[8] += "整体结构如图1所示。"
        self.assertIsNone(validate_visual_references(slots))

    def test_leading_reference_phrases_are_rejected(self) -> None:
        for phrase in (
            "如图1所示，", "见图1，", "参见图1，", "详见图1，", "图1展示了",
            "由图1可见，", "从图1可以看出，", "在图1中，", "根据图1，",
            "参考图1，", "结合图1，", "通过图1，", "（如图1所示）",
            '"如图1所示，', "　\t如 图 1 所示，", "\ufeff\u200b如图1所示，",
            "此外，如图1所示，", "同时，图1展示了",
        ):
            with self.subTest(phrase=phrase):
                slots = make_slots()
                slots[7] = phrase + "作者将特征处理分成多个阶段。"
                with self.assertRaisesRegex(ValueError, "must not lead paragraph"):
                    validate_visual_references(slots)

    def test_leading_table_reference_is_rejected(self) -> None:
        slots = make_slots()
        slots[20] = "如表1所示，作者进行了方法对比。"
        with self.assertRaisesRegex(ValueError, "must not lead paragraph"):
            validate_visual_references(slots)

    def test_opening_and_conclusion_have_the_same_paragraph_start_rule(self) -> None:
        for index in (3, 4, 24):
            with self.subTest(index=index):
                slots = make_slots()
                slots[index] = "如图1所示，作者构建了分层处理框架。"
                with self.assertRaisesRegex(ValueError, "must not lead paragraph"):
                    validate_visual_references(slots)

    def test_callout_at_start_of_later_sentence_is_allowed(self) -> None:
        slots = make_slots()
        slots[7] = "作者将特征处理分成多个阶段。如图1所示，不同阶段承担不同功能。"
        self.assertIsNone(validate_visual_references(slots))

    def test_grouped_references_ranges_panels_and_unicode_spacing(self) -> None:
        for text, expected in (
            ("结构见图2和图3。", {("图", 2), ("图", 3)}),
            ("结构见图2、3及4。", {("图", 2), ("图", 3), ("图", 4)}),
            ("结果见图2–图4。", {("图", 2), ("图", 3), ("图", 4)}),
            ("结果见图1至3与表1。", {("图", 1), ("图", 2), ("图", 3), ("表", 1)}),
            ("细节见图 ２(a)，定量结果见表 １。", {("图", 2), ("表", 1)}),
        ):
            with self.subTest(text=text):
                self.assertEqual(extract_visual_references(text), expected)

    def test_unknown_labels_are_rejected_without_single_digit_prefix_matches(self) -> None:
        for label in ("图10", "图0", "表2", "图1–100000000"):
            with self.subTest(label=label):
                slots = make_slots()
                slots[8] += f"另外的对比见{label}。"
                with self.assertRaisesRegex(ValueError, "Unknown local figure/table reference"):
                    validate_visual_references(slots)

    def test_descending_range_is_rejected(self) -> None:
        with self.assertRaisesRegex(ValueError, "ascending order"):
            extract_visual_references("结果见图4–2。")

    def test_ordinary_image_words_and_experimental_numbers_are_not_callouts(self) -> None:
        self.assertEqual(extract_visual_references("图像包含3个目标，表征维度为256。"), set())

    def test_quantity_after_a_reference_is_not_an_implicit_extra_figure(self) -> None:
        for quantity in ("2个子模块", "2次实验", "2.5倍增益", "2%提升"):
            with self.subTest(quantity=quantity):
                slots = make_slots()
                slots[11] = f"该模块的设计参考见图1，{quantity}用于检验配置差异。"
                with self.assertRaisesRegex(ValueError, "Missing body reference to 图2"):
                    validate_visual_references(slots)


class VisualReferenceIntegrationTests(unittest.TestCase):
    def test_manifest_loader_accepts_compliant_prose(self) -> None:
        with tempfile.TemporaryDirectory(prefix="compot-reference-test-") as directory:
            path = Path(directory) / "manifest.json"
            path.write_text(json.dumps(make_manifest(), ensure_ascii=False), encoding="utf-8")
            _, slots = load_manifest(path)
            self.assertEqual(slots[7], make_slots()[7])

    def test_manifest_loader_rejects_missing_or_leading_reference(self) -> None:
        for text, message in (
            ("该模块用于筛选有效信息。", "Missing body reference to 图3"),
            ("如图3所示，该模块筛选有效信息。", "must not lead paragraph"),
        ):
            with self.subTest(text=text), tempfile.TemporaryDirectory(prefix="compot-reference-test-") as directory:
                manifest = make_manifest()
                manifest["figure3"]["text"] = text
                path = Path(directory) / "manifest.json"
                path.write_text(json.dumps(manifest, ensure_ascii=False), encoding="utf-8")
                with self.assertRaisesRegex(ValueError, message):
                    load_manifest(path)

    def test_actual_saved_word_is_rechecked_after_manual_edit(self) -> None:
        with tempfile.TemporaryDirectory(prefix="compot-reference-test-") as directory:
            path = Path(directory) / "fixture.docx"
            document = make_document()
            document.save(path)
            loaded = Document(path)
            self.assertIsNone(validate_document_visual_references(loaded))
            loaded.paragraphs[23].text = "作者比较了不同方法的漏检情况。"
            loaded.save(path)
            with self.assertRaisesRegex(ValueError, "Missing body reference to 图5"):
                validate_document_visual_references(Document(path))

    def test_actual_word_leading_reference_is_rejected(self) -> None:
        document = make_document()
        document.paragraphs[14].text = "如图3所示，该模块筛选有效信息。"
        with self.assertRaisesRegex(ValueError, "must not lead paragraph"):
            validate_document_visual_references(document)


if __name__ == "__main__":
    unittest.main()
